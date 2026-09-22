#!/usr/bin/env python3
"""
migrate.py

Generalized version of migrate_comics.py: migrates one or more sheets
into ArchivesSpace as child archival objects of a resource, driven
entirely by a YAML mapping config (see mapping.py, configs/*.yaml,
and suggest_mapping.py for how to build one).

Usage examples:

  python migrate.py --config configs/peck_comics.yaml --xlsx Copy_of_Peck_Comics.xlsx --dry-run --limit 5
  python migrate.py --config configs/ciaraldi_sheet1.yaml --xlsx Ciaraldi_Standard_Size.xlsx --dry-run --limit 5

  # Override which sheet(s) the config applies to for this run:
  python migrate.py --config configs/ciaraldi_by_title.yaml --xlsx Ciaraldi_Standard_Size.xlsx \
      --sheets "Wizard,Flash" --dry-run

See README.md for the full flag list -- they're unchanged from
migrate_comics.py (--dry-run, --limit, --rows, --force, --publish,
--resource, --secrets, --state-file, --log-dir).

Row identity across MULTIPLE sheets: state.json keys are
"<sheet>::<excel row number>" rather than a bare row number, so
running several sheets in one config never collides.
"""

import argparse
import datetime
import json
import os
import sys

import openpyxl

from aspace_client import ArchivesSpaceClient, ArchivesSpaceError
from agents import resolve_publisher_agent
from people_agents import resolve_person_agent
from genres import resolve_genre_term
from digital_objects import resolve_digital_object, build_digital_object_instance
from containers import resolve_top_container_by_indicator
from hierarchy import HierarchyWalker
from locations import resolve_location, build_container_locations
from persistent_cache import load_cache, save_cache
from missing_locations import record_miss, write_missing_locations_report, write_pending_relink, coordinates_to_jsonable
from generic_builders import (
    build_title, build_title_and_digital_object, has_title, build_extents,
    build_notes, build_dates, parse_box_value, get_box_barcode_from_column, clean_str,
)
from mapping import MappingConfig, get_value
from resource_url import parse_target_reference, build_resource_ref, build_archival_object_ref, ResourceUrlError
from state import RunState

SECRETS_TEMPLATE = {
    "api_url": "",
    "repository_id": "",
    "username": "",
    "password": "",
}


def load_or_init_secrets(path: str) -> dict:
    if not os.path.exists(path):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(SECRETS_TEMPLATE, f, indent=2)
        print(f"No {path} found -- created a template for you.")
        print("Fill in api_url, repository_id, username, and password, then re-run.")
        sys.exit(1)
    with open(path, "r", encoding="utf-8") as f:
        secrets = json.load(f)
    missing = [k for k in SECRETS_TEMPLATE if not secrets.get(k)]
    if missing:
        print(f"{path} is missing values for: {', '.join(missing)}")
        sys.exit(1)
    return secrets


def read_sheet_rows(xlsx_path: str, sheet_name: str, has_header: bool, skip_rows=None):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name!r} not found. Available sheets: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return None, []

    if has_header:
        header = all_rows[0]
        header_index = {name: idx for idx, name in enumerate(header) if name}
        data_rows = list(enumerate(all_rows[1:], start=2))
    else:
        header_index = None
        data_rows = list(enumerate(all_rows, start=1))

    skip_rows = skip_rows or set()
    rows = [(n, r) for n, r in data_rows if any(c is not None for c in r) and n not in skip_rows]
    return header_index, rows


def _validate_mapping_against_header(mapping: MappingConfig, header_index, sheet_name: str):
    """Catch a real failure mode early: a name-based `column` in the
    config that doesn't actually exist in THIS sheet's header row.
    Without this check, get_value() would just silently return None
    for every row -- e.g. two sheets sharing a config where one
    spells a header "Starting Date" and another "Start Date" would
    quietly produce empty dates for the whole second sheet, with no
    error and no warning. Headerless sheets (header_index is None)
    skip this -- column refs there are positional by construction.
    """
    if header_index is None:
        return
    missing = []

    def check(column_ref, field_label):
        if isinstance(column_ref, str) and column_ref not in header_index:
            missing.append(f'{field_label} -> column "{column_ref}"')

    for field_label, cfg in (("title", mapping.title), ("publisher", mapping.publisher),
                              ("extent", mapping.extent), ("date", mapping.date),
                              ("physdesc", mapping.physdesc), ("box", mapping.box),
                              ("genre", mapping.genre), ("title_or_digital_object", mapping.title_or_digital_object),
                              ("ignore_column", mapping.ignore_column)):
        if cfg:
            check(cfg.get("column"), field_label)
    if mapping.box and mapping.box.get("barcode_column"):
        check(mapping.box.get("barcode_column"), "box.barcode_column")
    if mapping.location:
        for level in (1, 2, 3):
            level_cfg = mapping.location.get(f"coordinate_{level}")
            if level_cfg:
                check(level_cfg.get("column"), f"location.coordinate_{level}")
    for i, cfg in enumerate(mapping.scope_notes or []):
        check(cfg.get("column"), f"scope_notes[{i}]")
    for i, cfg in enumerate(mapping.agents or []):
        check(cfg.get("column"), f"agents[{i}]")
    for i, cfg in enumerate(mapping.hierarchy or []):
        check(cfg.get("id_column"), f"hierarchy[{i}].id_column")
        check(cfg.get("title_column"), f"hierarchy[{i}].title_column")

    if missing:
        actual = ", ".join(repr(h) for h in header_index if h)
        raise SystemExit(
            f'Config {mapping.path!r} references column(s) not present in sheet '
            f'{sheet_name!r}:\n  ' + "\n  ".join(missing) +
            f"\n\nThis sheet's actual headers: {actual}\n"
            f"(If sheets sharing this config have slightly different header text, "
            f"use column_index (0-based position) instead of column names.)"
        )


def build_archival_object(row_values, header_index, mapping: MappingConfig,
                           resource_ref, parent_ref, agent_links, genre_link,
                           container_link, digital_object_link, warnings):
    title, merge_url = build_title_and_digital_object(
        row_values, header_index, mapping.title, mapping.title_or_digital_object,
    )
    instances = _build_instances(container_link, mapping.instance_type)
    if digital_object_link:
        instances.append(build_digital_object_instance(digital_object_link))

    payload = {
        "jsonmodel_type": "archival_object",
        "title": title,
        "level": mapping.level,
        "publish": mapping.publish_default,
        "resource": {"ref": resource_ref},
        "dates": build_dates(row_values, header_index, mapping.date, warnings, title),
        "extents": build_extents(row_values, header_index, mapping.extent, warnings, title),
        "notes": build_notes(row_values, header_index, mapping, warnings),
        "instances": instances,
        "linked_agents": _build_linked_agents(agent_links),
        "subjects": _build_subjects(genre_link),
    }
    if parent_ref:
        payload["parent"] = {"ref": parent_ref}
    return payload, title, merge_url


def _build_instances(container_link, instance_type):
    if not container_link:
        return []
    return [{
        "jsonmodel_type": "instance",
        "instance_type": instance_type,
        "sub_container": {
            "jsonmodel_type": "sub_container",
            "top_container": {"ref": container_link["uri"]},
        },
    }]


def _build_linked_agents(agent_links):
    """agent_links: list of (agent_link_dict_or_None, role, relator)."""
    result = []
    for link, role, relator in agent_links or []:
        if link:
            result.append({"ref": link["uri"], "role": role, "relator": relator})
    return result


def _build_subjects(genre_link):
    if not genre_link:
        return []
    return [{"ref": genre_link["uri"]}]


def _build_coordinates(row_values, header_index, location_cfg):
    if not location_cfg:
        return {}
    coords = {}
    for level in (1, 2, 3):
        level_cfg = location_cfg.get(f"coordinate_{level}")
        if not level_cfg:
            continue
        raw = get_value(row_values, header_index, level_cfg.get("column"))
        if raw is None or str(raw).strip() == "":
            continue
        try:
            val = str(int(float(raw)))
        except (TypeError, ValueError):
            val = str(raw).strip()
        coords[level] = (level_cfg.get("label", f"Coordinate {level}"), val)
    return coords


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", required=True, help="Path to a mapping YAML config.")
    parser.add_argument("--xlsx", required=True, help="Path to the spreadsheet.")
    parser.add_argument("--sheets", default=None,
                         help="Comma-separated sheet names to process, overriding the config's own `sheets:` list.")
    parser.add_argument("--secrets", default="secrets.json")
    parser.add_argument("--resource", default=None)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N data rows PER SHEET.")
    parser.add_argument("--rows", default=None,
                         help="Comma-separated Excel row numbers to process (applies within each sheet processed).")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--state-file", default="state.json")
    parser.add_argument("--log-dir", default="logs")
    parser.add_argument("--pending-relink-file", default="pending_relink.json",
                         help="Where to accumulate top containers created without a location match, for relink_locations.py.")
    parser.add_argument("--container-cache-file", default="container_cache.json",
                         help="Remembers top containers this script has created, across separate runs -- "
                              "so re-running doesn't recreate ones ArchivesSpace's search can't reliably "
                              "find again (a real, confirmed issue on some instances; see KNOWN_LIMITATIONS.md).")
    parser.add_argument("--publish", action="store_true",
                         help="Override the config's publish_default to True.")
    args = parser.parse_args()

    mapping = MappingConfig.load(args.config)
    if args.publish:
        mapping.publish_default = True

    sheets = [s.strip() for s in args.sheets.split(",")] if args.sheets else mapping.sheets
    if not sheets:
        print("No sheets to process -- pass --sheets or set `sheets:` in the config.")
        sys.exit(1)

    secrets = load_or_init_secrets(args.secrets)

    target_input = args.resource or input(
        "Paste the target's staff URL, public URL, API URI, or numeric ID -- a "
        "RESOURCE (rows become direct children) or an ARCHIVAL OBJECT (rows nest "
        "under it) both work: "
    ).strip()
    try:
        target_kind, target_id, url_repo_id = parse_target_reference(target_input)
    except ResourceUrlError as exc:
        print(str(exc))
        sys.exit(1)

    repository_id = secrets["repository_id"]
    if url_repo_id and str(url_repo_id) != str(repository_id):
        print(f"NOTE: the URL you pasted references repository {url_repo_id}, "
              f"but secrets.json has repository_id={repository_id}. Using secrets.json's value.")

    os.makedirs(args.log_dir, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    log_path = os.path.join(args.log_dir, f"run-{timestamp}.log")
    log_file = open(log_path, "w", encoding="utf-8")

    def log(msg):
        print(msg)
        log_file.write(msg + "\n")
        log_file.flush()

    log(f"=== Run started {timestamp} ===")
    log(f"Config: {args.config} ({mapping.name})")
    log(f"Spreadsheet: {args.xlsx} | Sheets: {sheets}")
    log(f"Dry run: {args.dry_run} | Publish new archival objects: {mapping.publish_default}")

    try:
        client = ArchivesSpaceClient(
            api_url=secrets["api_url"], username=secrets["username"], password=secrets["password"],
            repository_id=repository_id, dry_run=args.dry_run, log=log,
        )
    except ArchivesSpaceError as exc:
        log(f"Could not connect/log in to ArchivesSpace: {exc}")
        sys.exit(1)

    # Every archival_object payload needs a "resource" ref regardless of
    # target kind -- ArchivesSpace requires it even for objects deep in
    # a tree. Targeting a resource directly: that's it, no "parent".
    # Targeting an archival object: rows nest under IT (initial_parent_ref),
    # but "resource" still has to be that object's own OWNING resource,
    # which only a real GET can tell us -- not something we can construct
    # from the pasted URL/ID alone.
    initial_parent_ref = None
    if target_kind == "resource":
        resource_ref = build_resource_ref(repository_id, target_id)
    else:
        initial_parent_ref = build_archival_object_ref(repository_id, target_id)
        try:
            target_record = client.get(initial_parent_ref)
        except ArchivesSpaceError as exc:
            log(f"Could not fetch archival object {initial_parent_ref}: {exc}")
            sys.exit(1)
        if not target_record:
            log(f"Archival object {initial_parent_ref} not found.")
            sys.exit(1)
        resource_ref = (target_record.get("resource") or {}).get("ref")
        if not resource_ref:
            log(f"Archival object {initial_parent_ref} has no owning resource on record -- "
                f"can't proceed (every archival_object needs one).")
            sys.exit(1)
        log(f"Target archival object {initial_parent_ref} belongs to resource {resource_ref}")

    log(f"Target resource: {resource_ref}"
        + (f" (rows nest under archival object {initial_parent_ref})" if initial_parent_ref else ""))

    state = RunState(args.state_file, log=log, dry_run=args.dry_run)
    agent_cache, genre_cache, digital_object_cache, location_cache = {}, {}, {}, {}
    container_cache = load_cache(args.container_cache_file)
    if container_cache:
        log(f"Loaded {len(container_cache)} previously-created top container(s) from "
            f"{args.container_cache_file} -- these won't be recreated even if ArchivesSpace's "
            f"own search can't find them.")
    missing_locations = {}
    pending_relink = []
    wanted_rows = None
    if args.rows:
        wanted_rows = {int(x.strip()) for x in args.rows.split(",") if x.strip()}

    succeeded, skipped, errored, header_only = 0, 0, 0, 0
    all_warnings = []

    for sheet_name in sheets:
        header_index, rows = read_sheet_rows(args.xlsx, sheet_name, mapping.has_header, mapping.skip_rows)
        _validate_mapping_against_header(mapping, header_index, sheet_name)
        if wanted_rows is not None:
            rows = [(n, r) for n, r in rows if n in wanted_rows]
        elif args.limit:
            rows = rows[: args.limit]

        walker = HierarchyWalker(
            mapping.hierarchy, resource_ref, client, log, initial_parent_ref=initial_parent_ref,
        ) if mapping.hierarchy else None

        for excel_row_num, row_values in rows:
            state_key = f"{sheet_name}::{excel_row_num}"

            # The hierarchy walk must run for EVERY row in order, even
            # ones already marked successful, or a resumed run loses
            # track of "current series" for the rows after them.
            if walker:
                walker.update(row_values, header_index)

            if mapping.ignore_column:
                ignore_raw = get_value(row_values, header_index, mapping.ignore_column.get("column"))
                if clean_str(ignore_raw):
                    log(f"{state_key}: 'Ignore' column set ({ignore_raw!r}) -- skipping row entirely.")
                    continue

            if not has_title(row_values, header_index, mapping.title):
                # A pure hierarchy-header row (e.g. a series boundary
                # with no file-level data of its own) -- state was
                # already updated above; nothing else to do.
                header_only += 1
                continue

            if not args.force and state.was_successful(state_key):
                log(f"{state_key}: already succeeded previously -- skipping. (use --force to redo)")
                skipped += 1
                continue

            parent_ref = walker.current_parent_ref() if walker else initial_parent_ref

            row_warnings = []
            try:
                agent_links = []
                if mapping.publisher:
                    pub_raw = get_value(row_values, header_index, mapping.publisher.get("column"))
                    link = resolve_publisher_agent(pub_raw, client, agent_cache, log) if pub_raw else None
                    agent_links.append((link, "creator", "pbl"))
                for agent_cfg in mapping.agents:
                    raw = get_value(row_values, header_index, agent_cfg.get("column"))
                    if not raw:
                        continue
                    resolver = resolve_person_agent if agent_cfg.get("agent_type") == "person" else resolve_publisher_agent
                    link = resolver(raw, client, agent_cache, log)
                    agent_links.append((link, agent_cfg.get("role", "creator"), agent_cfg.get("relator")))

                genre_link = None
                if mapping.genre:
                    genre_raw = get_value(row_values, header_index, mapping.genre.get("column"))
                    if genre_raw:
                        genre_link = resolve_genre_term(genre_raw, client, genre_cache, log, mapping.vocabulary_ref)

                container_link = None
                if mapping.box:
                    box_raw = get_value(row_values, header_index, mapping.box.get("column"))
                    indicator, barcode, container_type = parse_box_value(box_raw, mapping.box)
                    barcode = barcode or get_box_barcode_from_column(row_values, header_index, mapping.box)

                    container_locations_payload = None
                    coords = _build_coordinates(row_values, header_index, mapping.location)
                    building = mapping.location.get("building") if mapping.location else None
                    location_link = None
                    if coords:
                        location_link = resolve_location(coords, building, client, location_cache, log)
                        if location_link:
                            container_locations_payload = build_container_locations(location_link)

                    if indicator:
                        container_link = resolve_top_container_by_indicator(
                            indicator, client, container_cache, resource_ref, log,
                            barcode=barcode, container_type=container_type,
                            container_locations=container_locations_payload,
                        )
                        if container_link and container_link["status"] == "created":
                            save_cache(args.container_cache_file, container_cache, log)
                        if coords and not location_link:
                            title_guess = build_title(row_values, header_index, mapping.title)
                            record_miss(missing_locations, coords, building, title_guess)
                            if container_link and container_link["status"] == "created":
                                pending_relink.append({
                                    "top_container_uri": container_link["uri"],
                                    "building": building,
                                    "coordinates": coordinates_to_jsonable(coords),
                                })

                ao_payload, title, merge_url = build_archival_object(
                    row_values, header_index, mapping, resource_ref, parent_ref,
                    agent_links, genre_link, container_link, None, row_warnings,
                )

                if merge_url:
                    do_link = resolve_digital_object(merge_url, title, client, digital_object_cache, log)
                    ao_payload["instances"].append(build_digital_object_instance(do_link))

                result = client.post(f"{client.repo_prefix}/archival_objects", ao_payload)
                ao_uri = result.get("uri")

                for w in row_warnings:
                    log(f"  WARNING: {w}")
                all_warnings.extend(row_warnings)

                log(f"{state_key} ({title}): created {ao_uri}")
                state.record(state_key, "success", title=title, archival_object_uri=ao_uri, warnings=row_warnings)
                succeeded += 1

            except Exception as exc:  # noqa: BLE001
                title_guess = build_title(row_values, header_index, mapping.title)
                log(f"{state_key} ({title_guess}): ERROR -- {exc}")
                state.record(state_key, "error", title=title_guess, error=str(exc))
                errored += 1

    log("")
    log("=== Summary ===")
    log(f"Succeeded: {succeeded}")
    log(f"Skipped (already done): {skipped}")
    log(f"Header-only rows (no file-level object): {header_only}")
    log(f"Errored: {errored}")
    log(f"Total warnings: {len(all_warnings)}")
    log(f"Full log: {log_path}")
    log(f"State file: {args.state_file}")
    if errored:
        log("Re-run the same command to retry only the errored/unprocessed rows.")

    if missing_locations:
        report_path = os.path.join(args.log_dir, f"missing-locations-{timestamp}.xlsx")
        write_missing_locations_report(missing_locations, report_path)
        log(f"Missing locations report: {report_path} "
            f"({len(missing_locations)} distinct location(s) not found in ArchivesSpace)")
    if pending_relink:
        relink_path = args.pending_relink_file
        write_pending_relink(pending_relink, relink_path)
        log(f"Pending relink file: {relink_path} "
            f"({len(pending_relink)} top container(s) created without a location link)")
        log("Once the missing locations are created in ArchivesSpace, run relink_locations.py "
            f"--pending-relink-file {relink_path} to attach them.")

    log_file.close()


if __name__ == "__main__":
    main()
